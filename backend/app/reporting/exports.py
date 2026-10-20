"""Bounded, snapshot-consistent exports from the same authorized report services."""
import csv
from dataclasses import replace
from datetime import date, datetime
from enum import Enum
import io
import json
import re
from tempfile import SpooledTemporaryFile

from fastapi import HTTPException
from openpyxl import Workbook
from starlette.responses import StreamingResponse

from ..lifecycle.analytics import utc
from . import activity, inventory, work


MAX_EXPORT_ROWS = 20000
EXPORT_BATCH_SIZE = 1000
COLUMNS = {
    "tasks": [
        "id", "title", "task_group_id", "status", "project_name", "project_status",
        "assignee_name", "assignment_redacted", "created_at", "assigned_at", "due_at", "completed_at",
        "total_subtasks", "completed_subtasks", "overdue",
    ],
    "employees": [
        "employee_id", "name", "email", "status", "assigned_parent_tasks", "assigned_subtasks",
        "completed_parent_tasks", "in_progress_parent_tasks", "pending_parent_tasks",
        "blocked_parent_tasks", "overdue_parent_tasks", "completion_percentage",
        "project_count", "active_assignments", "completed_assignments",
        "recorded_effort_seconds",
    ],
    "projects": [
        "project_id", "name", "client_name", "status", "managers", "members", "manager_count", "member_count",
        "members_truncated",
        "start_date", "end_date", "total_tasks", "completed_tasks", "in_progress_tasks",
        "pending_tasks", "blocked_tasks", "overdue_tasks", "completion_percentage",
        "total_subtasks", "completed_subtasks",
        "recorded_effort_seconds",
    ],
    "statuses": ["status", "count", "percentage", "project_count", "employee_count", "overdue_tasks"],
    "activity": [
        "date", "tasks_created", "tasks_started", "tasks_completed", "tasks_reopened",
        "subtasks_completed", "work_sessions_started", "work_intervals_started",
        "recorded_effort_seconds", "unique_working_seconds",
    ],
    "time": [
        "id", "name", "recorded_effort_seconds", "unique_working_seconds", "interval_count",
        "running_interval_count", "paused_interval_count", "resumed_interval_count",
    ],
}


def safe_cell(value):
    """Keep untrusted text literal in both CSV and Excel, including leading controls."""
    if value is None:
        return ""
    if isinstance(value, Enum):
        value = value.value
    if isinstance(value, (datetime, date)):
        value = value.isoformat()
    elif isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if not isinstance(value, str):
        return value
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", value)
    if value.startswith(("\t", "\r", "\n")) or value.lstrip().startswith(("=", "+", "-", "@")):
        value = "'" + value
    # Excel's cell limit is 32,767 characters; every exported field is bounded.
    return value[:32767]


def _payload(db, user, filters, as_of, report, group_by, grouping):
    if report == "time":
        return work.time_report(db, user, filters, as_of=as_of, group_by=group_by)
    if report == "activity":
        body = activity.activity_report(db, user, filters, as_of=as_of, grouping=grouping)
        body.setdefault("total", len(body["items"]))
        return body
    return getattr(inventory, report)(db, user, filters, as_of)


def export_response(db, user, filters, as_of, *, report, format, group_by="employee", grouping="daily"):
    if report not in COLUMNS or format not in {"csv", "xlsx"}:
        raise HTTPException(422, "Unsupported report or export format")
    as_of = utc(as_of)
    # Interval unions and date buckets are computed once; SQL inventory pages use
    # bounded batches. No authorization or filters are bypassed for a download.
    size = MAX_EXPORT_ROWS + 1 if report in {"time", "activity"} else EXPORT_BATCH_SIZE
    export_filters = replace(filters, page=1, page_size=size)
    first = _payload(db, user, export_filters, as_of, report, group_by, grouping)
    total = first["total"]
    if total > MAX_EXPORT_ROWS:
        raise HTTPException(422, "Export exceeds 20,000 rows. Narrow the filters; no partial file returned.")
    meta = {
        "report": report, "generated_at": as_of.isoformat(), **first["meta"],
        "sort": filters.sort or {"tasks": "created_at", "statuses": "status", "activity": "date"}.get(report, "name"),
        "direction": filters.direction,
        "group_by": group_by if report == "time" else grouping if report == "activity" else None,
        "exported_rows": total, "export_scope": "all filtered rows, independent of the displayed page",
    }
    if "summary" in first:
        meta["summary"] = first["summary"]
    if "totals" in first:
        meta["totals"] = first["totals"]

    def rows():
        yield from first["items"]
        for offset in range(size, total, size):
            page = replace(export_filters, page=offset // size + 1)
            yield from _payload(db, user, page, as_of, report, group_by, grouping)["items"]

    columns = COLUMNS[report]
    output = SpooledTemporaryFile(max_size=2 * 1024 * 1024, mode="w+b")
    try:
        if format == "csv":
            wrapper = io.TextIOWrapper(output, encoding="utf-8-sig", newline="", write_through=True)
            writer = csv.writer(wrapper, lineterminator="\r\n")
            writer.writerow(["Zea.Play Reporting & Analytics"])
            for key, value in meta.items():
                writer.writerow([key, safe_cell(value)])
            writer.writerow([])
            writer.writerow(columns)
            for row in rows():
                writer.writerow([safe_cell(row.get(key)) for key in columns])
            wrapper.flush()
            wrapper.detach()
            media_type = "text/csv; charset=utf-8"
        else:
            book = Workbook(write_only=True)
            sheet = book.create_sheet("Report")
            sheet.freeze_panes = "A2"
            sheet.append(columns)
            for row in rows():
                sheet.append([safe_cell(row.get(key)) for key in columns])
            details = book.create_sheet("Metadata")
            details.append(["Property", "Value"])
            for key, value in meta.items():
                details.append([key, safe_cell(value)])
            book.save(output)
            book.close()
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        size_bytes = output.tell()
        output.seek(0)
    except BaseException:
        output.close()
        raise

    def content():
        try:
            while chunk := output.read(64 * 1024):
                yield chunk
        finally:
            output.close()

    period = f"{filters.from_date or 'all'}_to_{filters.to_date or 'present'}"
    filename = f"zea-play-{report}-{period}-{as_of.strftime('%Y%m%dT%H%M%SZ')}.{format}"
    return StreamingResponse(content(), media_type=media_type, headers={
        "Content-Disposition": f'attachment; filename="{filename}"',
        "Content-Length": str(size_bytes),
        "Cache-Control": "private, no-store",
        "X-Content-Type-Options": "nosniff",
    })
