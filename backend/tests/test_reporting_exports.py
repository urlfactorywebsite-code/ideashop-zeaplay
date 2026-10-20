"""Parse actual CSV and XLSX downloads and compare them with authorized report data."""
import csv
from datetime import date, datetime, timedelta, timezone
import io
import json

from openpyxl import load_workbook
import pytest

from backend.app.reporting import exports
from backend.tests.test_reporting_inventory import dataset, client, engine, auth, request  # noqa: F401


def download(client, user, report="tasks", format="csv", **params):
    return client.get("/api/reports/export", headers=auth(user),
                      params={"report": report, "format": format, **params})


def csv_rows(response, report):
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    start = next(index for index, row in enumerate(rows) if row == exports.COLUMNS[report])
    meta = {row[0]: row[1] for row in rows[1:start] if len(row) == 2}
    return meta, [dict(zip(rows[start], row)) for row in rows[start + 1:]]


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_download_matches_combined_filters_and_ignores_ui_page(dataset, client, format):
    params = {"from": "2026-10-19", "to": "2026-10-19", "timezone": "Asia/Kolkata",
              "project_id": dataset.project.id, "employee_id": dataset.alice.id,
              "status": "DONE", "search": "Shared", "sort": "title", "direction": "desc"}
    display = request(client, dataset.owner, "tasks", **params)
    response = download(client, dataset.owner, format=format, page=100, page_size=1, **params)
    assert response.status_code == 200, response.text
    assert response.headers["cache-control"] == "private, no-store"
    assert f".{format}" in response.headers["content-disposition"]
    if format == "csv":
        meta, rows = csv_rows(response, "tasks")
    else:
        book = load_workbook(io.BytesIO(response.content), read_only=True, data_only=False)
        cells = list(book["Report"].values)
        rows = [dict(zip(cells[0], row)) for row in cells[1:]]
        meta = dict(list(book["Metadata"].values)[1:])
        book.close()
    assert [row["id"] for row in rows] == [row["id"] for row in display["items"]]
    assert meta["timezone"] == "Asia/Kolkata"
    assert meta["from"] == meta["to"] == "2026-10-19"
    assert json.loads(meta["filters"])["employee_id"] == dataset.alice.id
    assert meta["generated_at"] and meta["date_semantics"] == "task.created_at"


@pytest.mark.parametrize("report", ["employees", "projects", "statuses", "activity", "time"])
def test_every_export_is_readable_and_authorized(dataset, client, report):
    response = download(client, dataset.owner, report=report, format="xlsx",
                        project_id=dataset.project.id, **{"from": "2026-10-19", "to": "2026-10-20"})
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content), read_only=True)
    rows = list(book["Report"].values)
    assert list(rows[0]) == exports.COLUMNS[report]
    assert "Metadata" in book.sheetnames
    book.close()
    response = download(client, dataset.manager, report=report, project_id=dataset.hidden.id)
    assert response.status_code == 404
    assert download(client, dataset.owner, report=report, employee_id=dataset.outsider.id).status_code == 404


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_formula_injection_stays_literal_and_sensitive_fields_absent(dataset, client, format):
    from backend.app import models
    dangerous = ' \t=HYPERLINK("https://example.invalid","unsafe")'
    with dataset.factory() as db:
        task = db.get(models.Task, dataset.first.id)
        task.title = dangerous
        db.commit()
    response = download(client, dataset.owner, format=format, employee_id=dataset.alice.id)
    assert response.status_code == 200, response.text
    if format == "csv":
        _, rows = csv_rows(response, "tasks")
        assert next(row["title"] for row in rows if row["id"] == dataset.first.id) == "'" + dangerous
    else:
        book = load_workbook(io.BytesIO(response.content))
        sheet = book["Report"]
        title_cells = [row[1] for row in list(sheet.rows)[1:]]
        assert any(cell.value == "'" + dangerous and cell.data_type != "f" for cell in title_cells)
        book.close()
    for forbidden in ["hashed_password", "attachments", "description", "correction_reason", "storage_key"]:
        assert forbidden not in exports.COLUMNS["tasks"]


def test_empty_export_and_limits_never_return_partial_file(dataset, client, monkeypatch):
    response = download(client, dataset.owner, search="nothing matches")
    assert response.status_code == 200
    assert csv_rows(response, "tasks")[1] == []
    monkeypatch.setattr(exports, "MAX_EXPORT_ROWS", 1)
    response = download(client, dataset.owner)
    assert response.status_code == 422 and "no partial" in response.json()["detail"].lower()
    assert client.get("/api/reports/export").status_code == 401
    assert download(client, dataset.alice, employee_id=dataset.bob.id).status_code == 404


def test_safe_cells_cover_controls_unicode_and_legitimate_numbers():
    for value in ["=1+1", "+cmd", "-cmd", "@SUM(1)", "\t=1", "\r=1", " \n@command"]:
        assert exports.safe_cell(value).startswith("'")
    assert exports.safe_cell(-12.5) == -12.5
    assert exports.safe_cell("ગુજરાતી") == "ગુજરાતી"
    assert exports.safe_cell("\x00literal") == "literal"
    assert exports.safe_cell("\ufffe=1+1\uffff") == "'=1+1"
    assert exports.safe_cell("\ud800literal") == "literal"
    assert exports.safe_cell(date(2026, 10, 20)) == "2026-10-20"
    assert exports.safe_cell(datetime(2026, 10, 20, tzinfo=timezone.utc)).endswith("+00:00")


def test_batched_export_keeps_order_filters_and_one_snapshot(dataset, client, monkeypatch):
    original = exports._payload
    snapshots = []
    def observed(db, user, filters, as_of, report, group_by, grouping):
        snapshots.append((as_of, filters.page, filters.project_id, filters.sort))
        return original(db, user, filters, as_of, report, group_by, grouping)
    monkeypatch.setattr(exports, "EXPORT_BATCH_SIZE", 1)
    monkeypatch.setattr(exports, "_payload", observed)
    params = {"project_id": dataset.project.id, "sort": "title", "direction": "desc"}
    display = request(client, dataset.owner, "tasks", **params)
    response = download(client, dataset.owner, page=8, **params)
    assert response.status_code == 200, response.text
    metadata, rows = csv_rows(response, "tasks")
    assert [row["id"] for row in rows] == [row["id"] for row in display["items"]]
    assert [page for _, page, _, _ in snapshots] == list(range(1, len(rows) + 1))
    assert len({as_of for as_of, *_ in snapshots}) == 1
    assert {(project, sort) for _, _, project, sort in snapshots} == {(dataset.project.id, "title")}
    assert metadata["sort"] == "title"
    assert int(response.headers["content-length"]) == len(response.content)


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_time_fractional_daily_export_matches_sorted_display(dataset, client, format):
    from backend.app import models as m
    with dataset.factory() as db:
        start = datetime(2026, 10, 19, 18, 29, 59, 500000, tzinfo=timezone.utc)
        db.add(m.TimeEntry(tenant_id=dataset.owner.tenant_id, user_id=dataset.alice.id,
                          task_id=dataset.first.id, started_at=start, stopped_at=start + timedelta(seconds=1),
                          duration_seconds=1, status=m.TimeEntryStatusEnum.STOPPED))
        db.commit()
    params = {"group_by": "daily", "from": "2026-10-19", "to": "2026-10-20",
              "timezone": "Asia/Kolkata", "project_id": dataset.project.id,
              "sort": "name", "direction": "desc"}
    display = request(client, dataset.owner, "time", **params)
    response = download(client, dataset.owner, report="time", format=format, **params)
    assert response.status_code == 200, response.text
    if format == "csv":
        meta, rows = csv_rows(response, "time")
    else:
        book = load_workbook(io.BytesIO(response.content), read_only=True)
        cells = list(book["Report"].values)
        rows = [dict(zip(cells[0], row)) for row in cells[1:]]
        meta = dict(list(book["Metadata"].values)[1:])
        book.close()
    assert [row["id"] for row in rows] == [row["id"] for row in display["items"]]
    assert [float(row["recorded_effort_seconds"]) for row in rows] == [.5, .5]
    assert json.loads(meta["summary"])["recorded_effort_seconds"] == 1


def test_employee_project_effort_columns_follow_inventory_population(dataset, client):
    from backend.app import models as m
    with dataset.factory() as db:
        for task, user, seconds in [(dataset.first, dataset.alice, 60), (dataset.first, dataset.bob, 30),
                                    (dataset.private, dataset.alice, 200)]:
            start = datetime(2026, 10, 19, 13, tzinfo=timezone.utc)
            db.add(m.TimeEntry(tenant_id=dataset.owner.tenant_id, user_id=user.id, task_id=task.id,
                              started_at=start, stopped_at=start + timedelta(seconds=seconds),
                              duration_seconds=seconds, status=m.TimeEntryStatusEnum.STOPPED))
        db.commit()
    params = {"project_id": dataset.project.id, "status": "DONE", "from": "2026-10-19",
              "to": "2026-10-19", "sort": "recorded_effort_seconds", "direction": "desc"}
    employees = request(client, dataset.owner, "employees", **params)
    assert [row["recorded_effort_seconds"] for row in employees["items"][:2]] == [60, 30]
    assert "actual worker" in employees["meta"]["recorded_effort_semantics"]
    projects = request(client, dataset.owner, "projects", **params)
    assert projects["items"][0]["recorded_effort_seconds"] == 90
    for report in ["employees", "projects"]:
        response = download(client, dataset.owner, report=report, **params)
        assert response.status_code == 200, response.text
        meta, rows = csv_rows(response, report)
        assert sum(float(row["recorded_effort_seconds"]) for row in rows) == 90
        assert "recorded_effort_semantics" in meta
    # A creation-date exclusion removes those intervals from inventory effort even
    # though the dedicated Time report can still show that historical work.
    with dataset.factory() as db:
        db.get(m.Task, dataset.first.id).created_at = datetime(2026, 10, 18)
        db.commit()
    assert request(client, dataset.owner, "projects", **params)["items"][0]["recorded_effort_seconds"] == 0
