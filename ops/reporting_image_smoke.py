"""Run inside a network-isolated, read-only backend image with /tmp writable."""
import io
import json
import sys

from fastapi import FastAPI
from fastapi.testclient import TestClient
import openpyxl

from app.main import app
from app.reporting.exports import safe_cell
from app.routers.reports import router


expected = {
    "/options", "/overview", "/employees", "/employees/{report_employee_id}",
    "/tasks", "/statuses", "/projects", "/projects/{report_project_id}",
    "/activity", "/time", "/time/sessions", "/export",
}
paths = {path.removeprefix("/api/reports") for path in app.openapi()["paths"]
         if path.startswith("/api/reports")}
assert expected == paths, paths

# The real production router and auth dependencies, without unrelated lifespan
# workers. A new in-memory engine is configured by the container command.
isolated = FastAPI()
isolated.include_router(router, prefix="/api")
with TestClient(isolated) as client:
    for suffix in expected:
        url = "/api/reports" + suffix.replace("{report_employee_id}", "untrusted")\
            .replace("{report_project_id}", "untrusted")
        assert client.get(url).status_code == 401, url

book = openpyxl.Workbook(write_only=True)
sheet = book.create_sheet("Export")
sheet.append([safe_cell("=1+1"), "ગુજરાતી", 0.5])
buffer = io.BytesIO()
book.save(buffer)
book.close()
parsed = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()))
assert parsed["Export"]["A1"].data_type == "s"
assert parsed["Export"]["A1"].value == "'=1+1"
assert parsed["Export"]["B1"].value == "ગુજરાતી"
assert parsed["Export"]["C1"].value == 0.5
parsed.close()

print(json.dumps({
    "result": "PASS", "python": sys.version.split()[0], "openpyxl": openpyxl.__version__,
    "registered_report_routes": len(paths), "unauthenticated_routes_rejected": len(expected),
    "xlsx_literal_unicode_fractional_roundtrip": "PASS",
    "scope": "network-isolated Linux image smoke; no deployed database or application startup",
}))
